"""Supervisor diary rules: semesters, counters, alerts, action tasks, Word export."""
from datetime import date, timedelta
from pathlib import Path

from hub.services import settings as st

FIELDS = ["meeting_date", "mode", "attendees", "agenda", "discussion", "decisions",
          "supervisor_feedback", "reflection", "next_meeting_date"]
MODES = {"in_person": "In person", "online": "Online", "email": "E-mail"}


def semester_ranges(s: dict) -> dict:
    ranges = {}
    for n in (1, 2):
        a, b = st.as_date(s, f"semester{n}_start"), st.as_date(s, f"semester{n}_end")
        if a and b and a <= b:
            ranges[n] = (a, b)
    return ranges


def semester_for(d: date, ranges: dict):
    """Return (semester, warning). Outside both ranges -> nearest semester with a warning."""
    if not ranges:
        return None, "Semester dates are not set in Settings, so this entry has no semester."
    for n, (a, b) in ranges.items():
        if a <= d <= b:
            return n, None

    def distance(n):
        a, b = ranges[n]
        return (a - d).days if d < a else (d - b).days

    n = min(ranges, key=distance)
    return n, f"{d.isoformat()} is outside both semesters; counted in Semester {n} (nearest)."


def current_semester(today: date, ranges: dict):
    for n, (a, b) in ranges.items():
        if a <= today <= b:
            return n
    return None


def counts(entries, semester) -> dict:
    in_sem = [e for e in entries if e["semester"] == semester]
    return {"signed": sum(1 for e in in_sem if e["signed"]),
            "unsigned": sum(1 for e in in_sem if not e["signed"]),
            "total": len(in_sem)}


def alerts(today: date, entries, s: dict) -> list[dict]:
    """Diary alerts for the dashboard. entries: rows with meeting_date, next_meeting_date, signed, semester."""
    out = []
    gap_days = st.as_int(s, "meeting_gap_days")
    required = st.as_int(s, "signed_meetings_per_semester")
    ranges = semester_ranges(s)
    today_s = today.isoformat()

    past = [e["meeting_date"] for e in entries if e["meeting_date"] <= today_s]
    has_future_plan = any((e["next_meeting_date"] or "") >= today_s for e in entries if e["next_meeting_date"]) \
        or any(e["meeting_date"] > today_s for e in entries)
    if past:
        last = date.fromisoformat(max(past))
        gap = (today - last).days
        if gap > gap_days and not has_future_plan:
            out.append({"level": "danger", "text": f"No supervisor meeting for {gap} days (limit {gap_days}). Book the next one.",
                        "href": "/diary/new"})
    else:
        start = st.as_date(s, "project_start_date")
        if start and (today - start).days > gap_days and not has_future_plan:
            out.append({"level": "danger", "text": "No supervisor meetings recorded yet.", "href": "/diary/new"})

    sem = current_semester(today, ranges)
    if sem:
        a, b = ranges[sem]
        span = max((b - a).days, 1)
        elapsed = (today - a).days / span
        signed = counts(entries, sem)["signed"]
        if (elapsed > 0.8 and signed < required) or (elapsed > 0.5 and signed < 3):
            out.append({"level": "danger",
                        "text": f"Signed meetings behind pace: {signed}/{required} with {round(elapsed * 100)}% of Semester {sem} gone.",
                        "href": "/diary"})

    stale = [e for e in entries if not e["signed"] and e["meeting_date"] <= (today - timedelta(days=7)).isoformat()]
    if stale:
        out.append({"level": "warning", "text": f"{len(stale)} diary {'entry' if len(stale) == 1 else 'entries'} older than 7 days still unsigned. Get them signed.",
                    "href": "/diary?signed=0"})
    return out


def sync_actions(conn, entry_id: int, actions) -> int:
    """Create a todo task per action row. Idempotent: matches existing tasks by (entry, title)."""
    created = 0
    existing = {r["title"].strip().lower() for r in
                conn.execute("SELECT title FROM tasks WHERE diary_entry_id = ?", (entry_id,))}
    for title, due in actions:
        title = (title or "").strip()
        if not title or title.lower() in existing:
            continue
        pos = conn.execute("SELECT COALESCE(MAX(position) + 1, 0) FROM tasks WHERE status = 'todo'").fetchone()[0]
        conn.execute("INSERT INTO tasks (title, status, due_date, diary_entry_id, position) VALUES (?, 'todo', ?, ?, ?)",
                     (title, due or None, entry_id, pos))
        existing.add(title.lower())
        created += 1
    return created


def save_entry(conn, data: dict, actions, s: dict, entry_id: int | None = None):
    """Insert or update an entry, derive its semester and sync action tasks. Returns (id, warning, created)."""
    sem, warning = semester_for(date.fromisoformat(data["meeting_date"]), semester_ranges(s))
    values = [data.get(f) or None for f in FIELDS]
    with conn:
        if entry_id is None:
            cur = conn.execute(
                f"INSERT INTO diary_entries ({', '.join(FIELDS)}, semester) VALUES ({', '.join('?' * len(FIELDS))}, ?)",
                values + [sem])
            entry_id = cur.lastrowid
        else:
            conn.execute(
                f"UPDATE diary_entries SET {', '.join(f + ' = ?' for f in FIELDS)}, semester = ?, "
                "updated_at = datetime('now','localtime') WHERE id = ?", values + [sem, entry_id])
        created = sync_actions(conn, entry_id, actions)
    return entry_id, warning, created


# ---- Word export -------------------------------------------------

def _doc(s: dict):
    from docx import Document
    from docx.shared import Cm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.2))
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    header = sec.header.paragraphs[0]
    header.text = (f"{s['project_title']}\n{s['student_name']} ({s['reg_no']}) · Supervisor: {s['supervisor_name']}")
    header.style.font.size = Pt(8.5)
    return doc


def _entry(doc, e, actions):
    doc.add_heading(f"Meeting on {e['meeting_date']}", level=2)
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    rows = [("Date", e["meeting_date"]), ("Mode", MODES.get(e["mode"] or "", e["mode"] or "")),
            ("Attendees", e["attendees"]), ("Agenda / questions", e["agenda"]),
            ("Discussion summary", e["discussion"]), ("Decisions made", e["decisions"]),
            ("Supervisor feedback", e["supervisor_feedback"]), ("Reflection", e["reflection"]),
            ("Next meeting", e["next_meeting_date"])]
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[0].paragraphs[0].runs[0].bold = True
        cells[1].text = value or ""
    doc.add_paragraph()
    doc.add_paragraph("Actions before next meeting").runs[0].bold = True
    if actions:
        for a in actions:
            doc.add_paragraph(a["title"] + (f" (due {a['due_date']})" if a["due_date"] else ""), style="List Number")
    else:
        doc.add_paragraph("None recorded.")
    doc.add_paragraph()
    signed = f"Signed on {e['signed_date']}" if e["signed"] and e["signed_date"] else ""
    doc.add_paragraph("Supervisor signature: ______________________    Date: ______________")
    if signed:
        doc.add_paragraph(signed).runs[0].italic = True


def _actions(conn, entry_id):
    return conn.execute("SELECT title, due_date FROM tasks WHERE diary_entry_id = ? ORDER BY id", (entry_id,)).fetchall()


def export_entry(conn, entry, s: dict, out_dir: Path, today: date) -> Path:
    doc = _doc(s)
    _entry(doc, entry, _actions(conn, entry["id"]))
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"diary_{entry['meeting_date']}_{entry['id']}_{today.isoformat()}.docx"
    doc.save(path)
    return path


def export_semester(conn, semester: int, s: dict, out_dir: Path, today: date) -> Path:
    from docx.enum.text import WD_BREAK

    entries = conn.execute("SELECT * FROM diary_entries WHERE semester = ? ORDER BY meeting_date", (semester,)).fetchall()
    c = counts(entries, semester)
    doc = _doc(s)
    doc.add_heading("Project diary", level=0)
    doc.add_paragraph(s["project_title"])
    doc.add_paragraph(f"Student: {s['student_name']} ({s['reg_no']})\nSupervisor: {s['supervisor_name']}\nSemester {semester}")
    doc.add_paragraph(f"Meetings recorded: {c['total']}\nSigned: {c['signed']} of {st.as_int(s, 'signed_meetings_per_semester')} required\n"
                      f"Unsigned: {c['unsigned']}\nExported: {today.isoformat()}")
    for e in entries:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _entry(doc, e, _actions(conn, e["id"]))
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"diary_S{semester}_{today.isoformat()}.docx"
    doc.save(path)
    return path
