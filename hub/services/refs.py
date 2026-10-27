"""Reference library rules: save, cite keys, merge, numbering, exports."""
import csv
import json
import shutil
from pathlib import Path

from hub.services import bibtex, ieee, refs_text

TYPES = {"journal": "Journal article", "conference": "Conference paper", "preprint": "Preprint", "book": "Book",
         "chapter": "Book chapter", "thesis": "Thesis", "report": "Report", "website": "Website",
         "dataset": "Dataset", "software": "Software"}
STATUSES = {"to_read": "To read", "reading": "Reading", "read": "Read", "cited": "Cited", "excluded": "Excluded"}
MATRIX = [("lm_country", "Country"), ("lm_dataset", "Dataset"), ("lm_sample", "Sample"),
          ("lm_label_source", "Label source"), ("lm_method", "Method"), ("lm_metrics", "Metrics"),
          ("lm_key_result", "Key result"), ("lm_limitations", "Limitations"), ("lm_relevance", "Relevance")]
FIELDS = ["ref_type", "title", "authors_json", "year", "month", "container", "volume", "issue", "pages", "article_no",
          "publisher", "location", "doi", "url", "arxiv_id", "accessed_date", "abstract", "pdf_path", "status",
          "rating", "notes", "ieee_number", "chapter_ids"] + [k for k, _ in MATRIX]
MERGE_FIELDS = [f for f in FIELDS if f not in ("notes",)]


def unique_cite_key(conn, base: str, exclude_id: int | None = None) -> str:
    taken = {r[0] for r in conn.execute("SELECT cite_key FROM refs WHERE cite_key LIKE ? AND id IS NOT ?",
                                        (base + "%", exclude_id))}
    if base not in taken:
        return base
    for suffix in "abcdefghijklmnopqrstuvwxyz":
        if base + suffix not in taken:
            return base + suffix
    n = 2
    while f"{base}{n}" in taken:
        n += 1
    return f"{base}{n}"


def clean(data: dict) -> dict:
    """Normalise a dict of form/lookup values into column values."""
    out = {}
    for f in FIELDS:
        v = data.get(f)
        if isinstance(v, str):
            v = v.strip()
        out[f] = v if v not in ("", None) else None
    if "authors" in data and data.get("authors") is not None:
        out["authors_json"] = json.dumps([{"given": (a.get("given") or "").strip(), "family": (a.get("family") or "").strip()}
                                          for a in data["authors"] if (a.get("family") or a.get("given") or "").strip()],
                                         ensure_ascii=False)
    out["authors_json"] = out["authors_json"] or "[]"
    chapters = data.get("chapter_ids")
    if isinstance(chapters, (list, tuple)):
        ids = sorted({int(c) for c in chapters if str(c).isdigit()})
        out["chapter_ids"] = json.dumps(ids) if ids else None
    for f in ("year", "month", "rating", "ieee_number"):
        try:
            out[f] = int(out[f]) if out[f] is not None else None
        except (TypeError, ValueError):
            out[f] = None
    if out["doi"]:
        from hub.services.crossref import normalise_doi
        out["doi"] = normalise_doi(out["doi"])
    out["ref_type"] = out["ref_type"] if out["ref_type"] in TYPES else "journal"
    out["status"] = out["status"] if out["status"] in STATUSES else "to_read"
    return out


def authors_text(authors_json: str) -> str:
    """'Family, Given' per line, for the edit form."""
    return "\n".join(f"{a.get('family', '')}, {a.get('given', '')}".strip(", ") for a in json.loads(authors_json or "[]"))


def parse_authors_text(text: str) -> list[dict]:
    out = []
    for line in (text or "").splitlines():
        line = line.strip()
        if not line:
            continue
        if "," in line:
            family, given = [p.strip() for p in line.split(",", 1)]
        else:
            given, _, family = line.rpartition(" ")
        out.append({"given": given, "family": family})
    return out


def save(conn, data: dict, ref_id: int | None = None) -> int:
    values = clean(data)
    authors = json.loads(values["authors_json"])
    with conn:
        if ref_id is None:
            key = unique_cite_key(conn, refs_text.cite_key_base(authors, values["year"], values["title"]))
            cur = conn.execute(f"INSERT INTO refs ({', '.join(FIELDS)}, cite_key) VALUES ({', '.join('?' * len(FIELDS))}, ?)",
                               [values[f] for f in FIELDS] + [key])
            return cur.lastrowid
        conn.execute(f"UPDATE refs SET {', '.join(f + ' = ?' for f in FIELDS)}, updated_at = datetime('now','localtime') "
                     "WHERE id = ?", [values[f] for f in FIELDS] + [ref_id])
    return ref_id


def update_fields(conn, ref_id: int, fields: dict):
    allowed = {k: v for k, v in fields.items() if k in FIELDS}
    if not allowed:
        return
    with conn:
        conn.execute(f"UPDATE refs SET {', '.join(k + ' = ?' for k in allowed)}, updated_at = datetime('now','localtime') "
                     "WHERE id = ?", [v if v != "" else None for v in allowed.values()] + [ref_id])


def merge(conn, older_id: int, newer_id: int):
    """Keep the older record, copy the newer one's non-empty fields into it, re-point links, delete the newer."""
    older = conn.execute("SELECT * FROM refs WHERE id = ?", (older_id,)).fetchone()
    newer = conn.execute("SELECT * FROM refs WHERE id = ?", (newer_id,)).fetchone()
    if not older or not newer or older_id == newer_id:
        raise ValueError("Both references must exist and differ.")
    updates = {f: newer[f] for f in MERGE_FIELDS
               if newer[f] not in (None, "", "[]") and not (f == "status" and newer[f] == "to_read")}
    notes = "\n\n".join(n for n in (older["notes"], newer["notes"]) if n)
    with conn:
        doi = updates.pop("doi", None)
        conn.execute("UPDATE refs SET doi = NULL WHERE id = ?", (newer_id,))  # free the unique DOI first
        if doi:
            updates["doi"] = doi
        updates["notes"] = notes or None
        conn.execute(f"UPDATE refs SET {', '.join(k + ' = ?' for k in updates)}, updated_at = datetime('now','localtime') "
                     "WHERE id = ?", list(updates.values()) + [older_id])
        conn.execute("UPDATE highlights SET ref_id = ? WHERE ref_id = ?", (older_id, newer_id))
        conn.execute("UPDATE OR IGNORE taggings SET item_id = ? WHERE item_type = 'ref' AND item_id = ?", (older_id, newer_id))
        conn.execute("DELETE FROM taggings WHERE item_type = 'ref' AND item_id = ?", (newer_id,))
        conn.execute("UPDATE OR IGNORE links SET from_id = ? WHERE from_type = 'ref' AND from_id = ?", (older_id, newer_id))
        conn.execute("UPDATE OR IGNORE links SET to_id = ? WHERE to_type = 'ref' AND to_id = ?", (older_id, newer_id))
        conn.execute("DELETE FROM links WHERE (from_type = 'ref' AND from_id = ?) OR (to_type = 'ref' AND to_id = ?)",
                     (newer_id, newer_id))
        conn.execute("DELETE FROM refs WHERE id = ?", (newer_id,))


def delete(conn, ref_id: int):
    with conn:
        conn.execute("DELETE FROM taggings WHERE item_type = 'ref' AND item_id = ?", (ref_id,))
        conn.execute("DELETE FROM links WHERE (from_type = 'ref' AND from_id = ?) OR (to_type = 'ref' AND to_id = ?)",
                     (ref_id, ref_id))
        conn.execute("DELETE FROM refs WHERE id = ?", (ref_id,))


def numbering(conn, status: str | None = "cited") -> list:
    """Ordered refs for the bibliography: ieee_number first (ascending), then first-added order."""
    sql = "SELECT * FROM refs"
    args = []
    if status:
        sql += " WHERE status = ?"
        args.append(status)
    sql += " ORDER BY CASE WHEN ieee_number IS NULL THEN 1 ELSE 0 END, ieee_number, id"
    return conn.execute(sql, args).fetchall()


def number_map(conn) -> dict[int, int]:
    return {r["id"]: i for i, r in enumerate(numbering(conn), start=1)}


# ---- links, tags ----------------------------------------------------------------

def link(conn, from_type, from_id, to_type, to_id):
    with conn:
        conn.execute("INSERT OR IGNORE INTO links (from_type, from_id, to_type, to_id) VALUES (?, ?, ?, ?)",
                     (from_type, from_id, to_type, to_id))


def unlink(conn, from_type, from_id, to_type, to_id):
    with conn:
        conn.execute("DELETE FROM links WHERE from_type=? AND from_id=? AND to_type=? AND to_id=?",
                     (from_type, from_id, to_type, to_id))


def linked(conn, item_type: str, item_id: int, other_type: str) -> list[int]:
    rows = conn.execute(
        "SELECT to_id AS id FROM links WHERE from_type=? AND from_id=? AND to_type=? "
        "UNION SELECT from_id FROM links WHERE to_type=? AND to_id=? AND from_type=?",
        (item_type, item_id, other_type, item_type, item_id, other_type)).fetchall()
    return [r["id"] for r in rows]


def set_tags(conn, item_type: str, item_id: int, names: list[str]):
    with conn:
        conn.execute("DELETE FROM taggings WHERE item_type = ? AND item_id = ?", (item_type, item_id))
        for name in {n.strip().lower() for n in names if n.strip()}:
            conn.execute("INSERT OR IGNORE INTO tags (name) VALUES (?)", (name,))
            tag_id = conn.execute("SELECT id FROM tags WHERE name = ?", (name,)).fetchone()[0]
            conn.execute("INSERT OR IGNORE INTO taggings (tag_id, item_type, item_id) VALUES (?, ?, ?)",
                         (tag_id, item_type, item_id))


def tags_for(conn, item_type: str) -> dict[int, list[str]]:
    out = {}
    for r in conn.execute("SELECT tg.item_id, t.name FROM taggings tg JOIN tags t ON t.id = tg.tag_id "
                          "WHERE tg.item_type = ? ORDER BY t.name", (item_type,)):
        out.setdefault(r["item_id"], []).append(r["name"])
    return out


# ---- PDF helpers ------------------------------------------------------------------

def pdf_lookup_hints(path: Path) -> dict:
    """DOI from the first two pages, else the PDF metadata title."""
    import pymupdf as fitz
    from hub.services.crossref import find_doi
    hints = {"doi": None, "title": None}
    with fitz.open(path) as doc:
        text = "".join(doc[i].get_text() for i in range(min(2, doc.page_count)))
        hints["doi"] = find_doi(text)
        title = (doc.metadata or {}).get("title") or ""
        if len(title.strip()) > 8:
            hints["title"] = title.strip()
    return hints


def store_upload(src: Path, library: Path, cite_key: str) -> str:
    """Copy an uploaded PDF into data/library/{cite_key}.pdf. Returns the relative path stored in pdf_path."""
    library.mkdir(parents=True, exist_ok=True)
    dest = library / f"{cite_key}.pdf"
    n = 2
    while dest.exists():
        dest = library / f"{cite_key}-{n}.pdf"
        n += 1
    shutil.copyfile(src, dest)
    return dest.name


# ---- exports ------------------------------------------------------------------------

def matrix_rows(conn, ids=None):
    rows = conn.execute("SELECT * FROM refs WHERE status != 'excluded' ORDER BY year DESC, id").fetchall()
    return [r for r in rows if ids is None or r["id"] in ids]


def export_matrix_csv(rows, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Cite key", "Authors", "Year", "Title"] + [label for _, label in MATRIX])
        for r in rows:
            w.writerow([r["cite_key"], ieee.author_list(ieee.authors_of(r)), r["year"] or "", r["title"]]
                       + [r[k] or "" for k, _ in MATRIX])
    return path


def export_matrix_xlsx(rows, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Literature matrix"
    ws.append(["Cite key", "Authors", "Year", "Title"] + [label for _, label in MATRIX])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r["cite_key"], ieee.author_list(ieee.authors_of(r)), r["year"], r["title"]] + [r[k] or "" for k, _ in MATRIX])
    widths = [18, 28, 7, 45] + [22] * len(MATRIX)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    wb.save(path)
    return path


def export_bibliography_docx(rows, path: Path) -> Path:
    from docx import Document
    from docx.shared import Cm, Pt
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading("References", level=1)
    for i, r in enumerate(rows, start=1):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent, pf.first_line_indent = Cm(0.9), Cm(-0.9)
        pf.tab_stops.add_tab_stop(Cm(0.9))
        pf.space_after = Pt(3)
        p.add_run(f"[{i}]\t")
        for text, italic in ieee.format(r):
            p.add_run(text).italic = italic
    doc.save(path)
    return path


def export_bib(rows, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(bibtex.dump(r) for r in rows), encoding="utf-8")
    return path

