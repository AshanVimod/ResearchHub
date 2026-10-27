"""Research-folder indexer. Read only: it never writes to, moves or deletes research files."""
import fnmatch
import json
import logging
import os
import re
import sys
import threading
from datetime import datetime, timedelta
from pathlib import Path

from hub import db
from hub.services import dedupe
from hub.services import settings as st

log = logging.getLogger(__name__)
_lock = threading.Lock()
PROGRESS_EVERY = 25
MISSING_PURGE_DAYS = 30
TEXT_EXTS = {".md", ".txt", ".py", ".csv", ".r", ".sql", ".json", ".tex", ".bib"}
EXTRACTABLE = {".pdf", ".docx", ".xlsx", ".ipynb"} | TEXT_EXTS


def excluded(name: str, globs: list[str]) -> bool:
    low = name.lower()
    return any(fnmatch.fnmatch(low, g.lower()) for g in globs)


def walk(root: Path, globs: list[str]):
    """Yield (path, rel_path, stat) for every file under root, skipping excluded names. Unicode-safe."""
    stack = [root]
    while stack:
        current = stack.pop()
        try:
            with os.scandir(current) as it:
                entries = list(it)
        except OSError as exc:
            log.warning("Cannot read %s: %s", current, exc)
            continue
        for e in entries:
            if excluded(e.name, globs):
                continue
            try:
                if e.is_dir(follow_symlinks=False):
                    stack.append(Path(e.path))
                elif e.is_file(follow_symlinks=False):
                    p = Path(e.path)
                    yield p, p.relative_to(root).as_posix(), e.stat()
            except OSError as exc:
                log.warning("Cannot stat %s: %s", e.path, exc)


def has_search_index(conn) -> bool:
    return bool(conn.execute("SELECT 1 FROM sqlite_master WHERE name = 'search_index'").fetchone())


def os_path(path: Path) -> Path:
    """Path usable for opening on Windows even beyond the 260-character MAX_PATH limit."""
    if sys.platform == "win32":
        s = str(Path(path).absolute())
        if len(s) >= 240 and not s.startswith("\\\\?\\"):
            return Path("\\\\?\\UNC\\" + s[2:] if s.startswith("\\\\") else "\\\\?\\" + s)
    return Path(path)


def file_ext(path: Path) -> str:
    """Real extensions only: 'Paper 2021.16' or 'Title. hal, hal-05375997' have none."""
    ext = path.suffix.lower()
    return ext if re.fullmatch(r"\.[a-z][a-z0-9]{0,5}", ext) else ""


def sniff_pdf(path: Path) -> bool:
    with open(os_path(path), "rb") as f:
        return f.read(5) == b"%PDF-"


def _page_count(path: Path):
    import pymupdf as fitz
    with fitz.open(os_path(path)) as doc:
        if doc.needs_pass:
            raise ValueError("PDF is encrypted")
        return doc.page_count


def scan(conn, s: dict, now: datetime | None = None, job_id: int | None = None) -> dict:
    """Walk every research root, upsert changed files, flag vanished ones. Returns job counters."""
    now = now or datetime.now()
    stamp = now.isoformat(timespec="seconds")
    globs = st.as_list(s, "index_exclude_globs")
    counts = {"scanned": 0, "added": 0, "updated": 0, "removed": 0, "errors": 0}
    if job_id is None:
        with conn:
            job_id = conn.execute("INSERT INTO index_jobs (started_at, state) VALUES (?, 'running')", (stamp,)).lastrowid

    def progress(state="running", finished=None):
        with conn:
            conn.execute("UPDATE index_jobs SET scanned=?, added=?, updated=?, removed=?, errors=?, state=?, finished_at=?, "
                         "updated_at=datetime('now','localtime') WHERE id=?",
                         (counts["scanned"], counts["added"], counts["updated"], counts["removed"], counts["errors"],
                          state, finished, job_id))

    for root_str in st.as_list(s, "research_roots"):
        root = Path(root_str)
        if not root.is_dir():
            counts["errors"] += 1
            log.warning("Research root not found: %s", root)
            continue
        existing = {r["rel_path"]: r for r in conn.execute("SELECT * FROM files WHERE root = ?", (root_str,))}
        seen = set()
        for path, rel, stat in walk(root, globs):
            seen.add(rel)
            counts["scanned"] += 1
            row = existing.get(rel)
            if row and row["text_status"] != "missing" and row["size_bytes"] == stat.st_size and row["mtime"] == stat.st_mtime:
                if counts["scanned"] % PROGRESS_EVERY == 0:
                    progress()
                continue
            ext = file_ext(path)
            status, error, pages, sha = "skipped", None, None, None
            try:
                if not ext and sniff_pdf(path):  # downloads that lost their .pdf
                    ext = ".pdf"
                status = "pending" if ext in EXTRACTABLE else "skipped"
                sha = dedupe.sha256_file(os_path(path))
                if ext == ".pdf":
                    pages = _page_count(path)
            except Exception as exc:  # unreadable, locked, encrypted or corrupt
                status, error = "error", str(exc)[:300]
                counts["errors"] += 1
            with conn:
                conn.execute(
                    "INSERT INTO files (root, rel_path, name, ext, size_bytes, mtime, sha256, page_count, text_status, error, indexed_at) "
                    "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?) ON CONFLICT(root, rel_path) DO UPDATE SET "
                    "name=excluded.name, ext=excluded.ext, size_bytes=excluded.size_bytes, mtime=excluded.mtime, "
                    "sha256=excluded.sha256, page_count=excluded.page_count, text_status=excluded.text_status, "
                    "error=excluded.error, indexed_at=excluded.indexed_at, updated_at=datetime('now','localtime')",
                    (root_str, rel, path.name, ext, stat.st_size, stat.st_mtime, sha, pages, status, error, stamp))
            counts["updated" if row else "added"] += 1
            if counts["scanned"] % PROGRESS_EVERY == 0:
                progress()
        gone = [r for rel, r in existing.items() if rel not in seen and r["text_status"] != "missing"]
        with conn:
            for r in gone:
                conn.execute("UPDATE files SET text_status = 'missing', updated_at = ? WHERE id = ?", (stamp, r["id"]))
                if has_search_index(conn):
                    conn.execute("DELETE FROM search_index WHERE item_type = 'file_page' AND item_id = ?", (r["id"],))
        counts["removed"] += len(gone)

    cutoff = (now - timedelta(days=MISSING_PURGE_DAYS)).isoformat(timespec="seconds")
    with conn:
        conn.execute("DELETE FROM files WHERE text_status = 'missing' AND updated_at < ?", (cutoff,))
    if has_search_index(conn):
        extract_pending(conn, s)
        embed_if_enabled(conn, s)
    progress("done", datetime.now().isoformat(timespec="seconds"))
    return dict(counts, job_id=job_id)


# ---- text extraction (feeds search_index once migration 003 exists) ----------------

def extract_text(path: Path, max_mb: int, ext: str | None = None) -> list[tuple[int | None, str]]:
    """Return [(page or None, text)] for a file. PDF yields one item per page."""
    ext = ext if ext is not None else file_ext(path)
    path = os_path(path)
    if ext == ".pdf":
        import pymupdf as fitz
        with fitz.open(path) as doc:
            if doc.needs_pass:
                raise ValueError("PDF is encrypted")
            return [(i + 1, doc[i].get_text()) for i in range(doc.page_count)]
    if ext == ".docx":
        from docx import Document
        d = Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.extend(c.text for c in row.cells)
        return [(None, "\n".join(p for p in parts if p))]
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            n = 0
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if v not in (None, ""):
                        parts.append(str(v))
                        n += 1
                        if n >= 2000:
                            break
                if n >= 2000:
                    break
        wb.close()
        return [(None, " ".join(parts))]
    if ext == ".ipynb":
        nb = json.loads(path.read_text(encoding="utf-8", errors="replace"))
        cells = ["".join(c.get("source", "")) if isinstance(c.get("source"), list) else c.get("source", "")
                 for c in nb.get("cells", []) if c.get("cell_type") in ("markdown", "code")]
        return [(None, "\n".join(cells))]
    if ext in TEXT_EXTS:
        if path.stat().st_size > max_mb * 1024 * 1024:
            raise OverflowError("larger than index_max_text_mb")
        return [(None, path.read_text(encoding="utf-8", errors="replace"))]
    return []


def extract_pending(conn, s: dict):
    max_mb = st.as_int(s, "index_max_text_mb")
    for r in conn.execute("SELECT * FROM files WHERE text_status = 'pending'").fetchall():
        path = Path(r["root"]) / r["rel_path"]
        try:
            chunks = extract_text(path, max_mb, r["ext"])
        except OverflowError:
            with conn:
                conn.execute("UPDATE files SET text_status='skipped', error=? WHERE id=?", ("too large for text", r["id"]))
            continue
        except Exception as exc:
            with conn:
                conn.execute("UPDATE files SET text_status='error', error=? WHERE id=?", (str(exc)[:300], r["id"]))
            continue
        with conn:
            conn.execute("DELETE FROM search_index WHERE item_type='file_page' AND item_id=?", (r["id"],))
            if conn.execute("SELECT 1 FROM sqlite_master WHERE name = 'chunks'").fetchone():
                conn.execute("DELETE FROM chunks WHERE file_id = ?", (r["id"],))  # re-embedded from the new text
            conn.executemany("INSERT INTO search_index (title, body, item_type, item_id, page) VALUES (?, ?, 'file_page', ?, ?)",
                             [(r["name"], text, r["id"], page) for page, text in chunks if text.strip()])
            conn.execute("UPDATE files SET text_status='ok', error=NULL WHERE id=?", (r["id"],))


def embed_if_enabled(conn, s: dict):
    """Phase 4 semantic search: embed new page text when the feature flag is on and the packages exist."""
    from hub.services import semantic
    if not semantic.enabled(s) or not semantic.has_table(conn) or not semantic.available()[0]:
        return
    try:
        n = semantic.embed_pending(conn)
        if n:
            log.info("Embedded %d file(s) for semantic search", n)
    except Exception:  # model download failed offline, etc.; keyword search is unaffected
        log.exception("Semantic embedding failed")


# ---- background worker ------------------------------------------------------------

def is_running() -> bool:
    return _lock.locked()


def start_background(db_path, logger=None) -> bool:
    """Run one scan in a daemon thread. Returns False if a scan is already running."""
    if not _lock.acquire(blocking=False):
        return False

    def run():
        conn = db.connect(db_path)
        try:
            scan(conn, st.get_all(conn))
        except Exception:
            (logger or log).exception("Index scan failed")
            with conn:
                conn.execute("UPDATE index_jobs SET state='error', finished_at=? WHERE state='running'",
                             (datetime.now().isoformat(timespec="seconds"),))
        finally:
            conn.close()
            _lock.release()

    threading.Thread(target=run, name="indexer", daemon=True).start()
    return True


def schedule(db_path, logger=None, first_delay=10, every_minutes=30):
    """Scan 10 s after start, then every 30 minutes while the app runs."""
    def loop():
        start_background(db_path, logger)
        t = threading.Timer(every_minutes * 60, loop)
        t.daemon = True
        t.start()
    t = threading.Timer(first_delay, loop)
    t.daemon = True
    t.start()


def latest_job(conn):
    return conn.execute("SELECT * FROM index_jobs ORDER BY id DESC LIMIT 1").fetchone()


def mark_stale_jobs(conn):
    """A job left 'running' by a previous process can never finish."""
    with conn:
        conn.execute("UPDATE index_jobs SET state='interrupted' WHERE state='running'")
