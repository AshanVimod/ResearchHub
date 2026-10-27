"""Thesis word counts, chapters and checklists."""
import json
import re
import shutil
import tempfile
from datetime import datetime
from pathlib import Path

WORD_RE = re.compile(r"[A-Za-z0-9][A-Za-z0-9'’\-]*")
NUMBERING_RE = re.compile(r"^\s*(chapter\s+)?([0-9]+|[ivxlc]+)?[\s.:\-–)]*", re.I)
MIN_MAIN_WORDS = 10_000
ABSTRACT = (250, 300)

# ordinal, title, heading_match, target words, guideline note
CHAPTER_SEED = [
    (0, "Front matter (title, declaration, abstract, acknowledgement, contents, lists)", "", None, "Abstract 250–300 words"),
    (1, "Introduction", "Introduction", 1200, "Problem / gap, aims, objectives, motivation"),
    (2, "Literature Review", "Literature Review", 2800, "Thorough, well-structured, properly cited"),
    (3, "Methods", "Methods|Methodology|Research Methodology", 2200, "Design, data, sampling, ethics, analysis"),
    (4, "Results", "Results", 1800, "Numbered figures and tables with clear titles"),
    (5, "Discussion", "Discussion", 1300, "Interpretation and limitations"),
    (6, "Conclusion", "Conclusion", 500, "Back to aims and objectives"),
    (7, "Recommendation", "Recommendation", 400, "Further research and actions"),
]
NOTES = {o: n for o, *_, n in CHAPTER_SEED}

CHECKLISTS = {
    "SUBMISSION": [
        "Draft sent to supervisor at least 1 week before the deadline", "All supervisor corrections made",
        "Plagiarism check done (record tool and %)", "Supervisor approval before final upload",
        "Main text at least 10,000 words", "Abstract 250–300 words",
        "Figures and tables numbered consecutively with titles", "IEEE reference style checked",
        "Examiner corrections done (post-viva)", "Supervisor signed the final copy",
    ],
    "PROPOSAL": ["At most 10 pages", "Title at most 10 words", "All required template sections", "IEEE references",
                 "10-minute presentation rehearsed"],
    "MIDTERM": ["Progress slides", "Results so far", "Updated timeline", "Diary signed"],
}
CHECKLIST_TITLES = {"SUBMISSION": "Submission", "PROPOSAL": "Proposal", "MIDTERM": "Mid-term review"}
STATUSES = {"not_started": "Not started", "drafting": "Drafting", "with_supervisor": "With supervisor",
            "revising": "Revising", "final": "Final"}


def seed(conn):
    with conn:
        if not conn.execute("SELECT COUNT(*) FROM chapters").fetchone()[0]:
            conn.executemany("INSERT INTO chapters (ordinal, title, heading_match, target_words) VALUES (?, ?, ?, ?)",
                             [(o, t, h, w) for o, t, h, w, _ in CHAPTER_SEED])
        for code, items in CHECKLISTS.items():
            if not conn.execute("SELECT 1 FROM checklist_items WHERE list_code = ?", (code,)).fetchone():
                conn.executemany("INSERT INTO checklist_items (list_code, ordinal, text) VALUES (?, ?, ?)",
                                 [(code, i, t) for i, t in enumerate(items)])


def count_words(text: str) -> int:
    return len(WORD_RE.findall(text or ""))


def normalise_heading(text: str) -> str:
    return re.sub(r"\s+", " ", NUMBERING_RE.sub("", text or "", count=1)).strip().lower()


def heading_matches(text: str, heading_match: str) -> bool:
    h = normalise_heading(text)
    return any(alt.strip() and h.startswith(alt.strip().lower()) for alt in (heading_match or "").split("|"))


def iter_blocks(doc):
    """Yield (style_name, text) for body paragraphs and table cells, in document order."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            yield (p.style.name if p.style is not None else ""), p.text
        elif tag == "tbl":
            for row in Table(child, doc).rows:
                seen = set()
                for cell in row.cells:
                    if id(cell._tc) in seen:  # merged cells repeat
                        continue
                    seen.add(id(cell._tc))
                    for p in cell.paragraphs:
                        yield "", p.text


def analyse(doc, chapters) -> dict:
    """Split the document at Heading 1 paragraphs that match a chapter and count words per chapter.

    Text before the first chapter counts as front matter (ordinal 0). A Heading 1 that matches no
    chapter (References, Appendix A, ...) ends counting until the next matching chapter.
    """
    words = {c["id"]: 0 for c in chapters}
    front = next((c["id"] for c in chapters if c["ordinal"] == 0), None)
    current, started = front, False
    abstract_words, in_abstract, abstract_found = 0, False, False
    for style, text in iter_blocks(doc):
        style_l = style.lower()
        if style_l.startswith("heading") or style_l == "title":
            in_abstract = normalise_heading(text).startswith("abstract")
            abstract_found = abstract_found or in_abstract
        if style_l.startswith("heading 1"):
            match = next((c for c in chapters if c["ordinal"] and heading_matches(text, c["heading_match"])), None)
            if match:
                current, started = match["id"], True
            else:
                # Before chapter 1 an unmatched H1 (Declaration, Contents...) is still front matter;
                # after it (References, Appendix A...) counting stops.
                current = None if started else front
            continue
        n = count_words(text)
        if in_abstract:
            abstract_words += n
        if current is not None:
            words[current] += n
    main = sum(words[c["id"]] for c in chapters if 1 <= c["ordinal"] <= 7)
    return {"per_chapter": words, "main": main, "abstract": abstract_words if abstract_found else None}


def open_docx(path: Path):
    """Open the thesis; if Word has it locked, read a temporary copy."""
    from docx import Document
    try:
        return Document(str(path))
    except PermissionError:
        with tempfile.TemporaryDirectory() as tmp:
            copy = Path(tmp) / "thesis_copy.docx"
            shutil.copyfile(path, copy)
            return Document(str(copy))


def last_count(conn):
    return conn.execute("SELECT * FROM thesis_counts ORDER BY id DESC LIMIT 1").fetchone()


def recount(conn, path: Path, now: datetime, force: bool = False) -> dict:
    """Recount when the file's mtime changed (or force). Returns {'status': ..., 'message': ...}."""
    if not path or not str(path).strip():
        return {"status": "unset", "message": "Set the thesis .docx path in Settings to count words."}
    path = Path(path)
    if not path.exists():
        return {"status": "missing", "message": f"Thesis file not found: {path}"}
    mtime = path.stat().st_mtime
    last = last_count(conn)
    if last and last["file_mtime"] == mtime and not force:
        return {"status": "unchanged", "message": None}
    try:
        doc = open_docx(path)
    except Exception as exc:  # corrupt or unreadable
        return {"status": "error", "message": f"Could not read the thesis: {exc}"}
    chapters = conn.execute("SELECT * FROM chapters ORDER BY ordinal").fetchall()
    res = analyse(doc, chapters)
    stamp = now.isoformat(timespec="seconds")
    with conn:
        for cid, n in res["per_chapter"].items():
            conn.execute("UPDATE chapters SET current_words = ?, last_counted_at = ? WHERE id = ?", (n, stamp, cid))
        conn.execute("INSERT INTO thesis_counts (counted_at, file_mtime, main_words, abstract_words, per_chapter_json) "
                     "VALUES (?, ?, ?, ?, ?)", (stamp, mtime, res["main"], res["abstract"],
                                                json.dumps({str(k): v for k, v in res["per_chapter"].items()})))
    return {"status": "counted", "message": None}


def abstract_badge(n) -> str:
    if n is None:
        return "none"
    lo, hi = ABSTRACT
    if lo <= n <= hi:
        return "success"
    return "warning" if lo - 50 <= n <= hi + 50 else "danger"


def history(conn, limit: int = 120) -> list:
    rows = conn.execute("SELECT counted_at, main_words FROM thesis_counts ORDER BY id DESC LIMIT ?", (limit,)).fetchall()
    # Keep the last count of each day for the chart.
    by_day = {}
    for r in reversed(rows):
        by_day[r["counted_at"][:10]] = r["main_words"]
    return [{"day": d, "words": w} for d, w in by_day.items()]
